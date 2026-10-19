"""
Helpers shared by the tests.

valid/ holds one valid dataset, one CSV file per table. Each test works on a copy of it, and a broken dataset
is that copy with one thing changed, so that it breaks exactly one rule.
"""
import csv
import io
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from validation import DATE, Table, ValidationError, cell_text, load_schema, validate_files

ROOT = Path(__file__).resolve().parent.parent
SCHEMA = load_schema()
VALID = Path(__file__).resolve().parent / "valid"
TEMPLATE = ROOT / "template" / "collection-sheet.xlsx"

# Read once rather than for every test, which is slow when the repository is on a network drive
VALID_FILES = {path.name: path.read_bytes() for path in VALID.iterdir()}
TEMPLATE_BYTES = TEMPLATE.read_bytes()


class Dataset:
    """A copy of the valid dataset in a folder, with ways to change its CSV files."""

    def __init__(self, folder: Path):
        folder.mkdir()
        for name, content in VALID_FILES.items():
            (folder / name).write_bytes(content)
        self.folder = folder

    def path(self, table: str) -> Path:
        return self.folder / f"{table}.csv"

    def files(self) -> list[tuple[str, Path]]:
        """(file name, path) of each file in the folder, sorted by name, as validate_files takes them."""
        return [(path.name, path) for path in sorted(self.folder.iterdir(), key=lambda path: path.name)]

    def read(self, table: str) -> list[list[str]]:
        """The rows of a table's file, header first."""
        with self.path(table).open(encoding="utf-8", newline="") as f:
            return list(csv.reader(f))

    def write(self, table: str, rows: list[list[str]], delimiter: str = ",", encoding: str = "utf-8",
              line_end: str = "\n") -> None:
        with self.path(table).open("w", encoding=encoding, newline="") as f:
            csv.writer(f, delimiter=delimiter, lineterminator=line_end).writerows(rows)

    def ids(self, table: str) -> list[str]:
        """The value in the first column of each row below the header."""
        return [row[0] if row else "" for row in self.read(table)[1:]]

    def row(self, table: str, rid: str) -> int:
        """The row number in the file of the row with this ID. The header is row 1."""
        return self.ids(table).index(rid) + 2

    def where(self, table: str, rid: str, column: str) -> str:
        """Where a cell is, as the validator's messages say it, e.g. "Phenotype.csv, row 2, column Score"."""
        return f"{table}.csv, row {self.row(table, rid)}, column {column}"

    def get(self, table: str, rid: str, column: str) -> str:
        rows = self.read(table)
        return rows[self.row(table, rid) - 1][rows[0].index(column)]

    def set(self, table: str, rid: str, column: str, value: str) -> None:
        """Sets the cell in the given column of the row with this ID."""
        rows = self.read(table)
        rows[self.row(table, rid) - 1][rows[0].index(column)] = value
        self.write(table, rows)

    def add_row(self, table: str, **values: str) -> None:
        """Adds a row at the end, with the given values and the other cells empty."""
        rows = self.read(table)
        rows.append([values.get(column, "") for column in rows[0]])
        self.write(table, rows)

    def delete_row(self, table: str, rid: str) -> None:
        rows = self.read(table)
        del rows[self.row(table, rid) - 1]
        self.write(table, rows)

    def add_column(self, table: str, name: str, values: tuple[str, ...] = ()) -> None:
        """Adds a column at the end, with the given values in the first rows and the other cells empty."""
        rows = self.read(table)
        rows[0].append(name)
        for i, row in enumerate(rows[1:]):
            row.append(values[i] if i < len(values) else "")
        self.write(table, rows)

    def drop_column(self, table: str, column: str) -> None:
        rows = self.read(table)
        j = rows[0].index(column)
        self.write(table, [row[:j] + row[j + 1:] for row in rows])

    def rename_column(self, table: str, old: str, new: str) -> None:
        rows = self.read(table)
        rows[0][rows[0].index(old)] = new
        self.write(table, rows)

    def rename_file(self, table: str, name: str) -> None:
        self.path(table).rename(self.folder / name)

    def drop_table(self, table: str) -> None:
        self.path(table).unlink()


def excel_value(text: str, column_type: str):
    """
    The value someone filling in the collection sheet by hand would leave in the cell: a date as an Excel date,
    and a number as a number, as long as the validator reads it back as the same text.
    """
    if column_type == "date" and DATE.fullmatch(text):
        try:
            return datetime.fromisoformat(text)
        except ValueError:
            return text
    if column_type in ("integer", "number"):
        for convert in (int, float):
            try:
                value = convert(text)
            except ValueError:
                continue
            if cell_text(value) == text:
                return value
    return text


def make_workbook(dataset: Dataset, path: Path, descriptor_row: bool = True) -> Path:
    """
    Fills a copy of the blank collection sheet with the dataset's tables, from row 3 (below the row of column
    descriptions), or from row 2 with the descriptions removed. Tables the dataset lacks keep their blank sheet.
    """
    wb = openpyxl.load_workbook(io.BytesIO(TEMPLATE_BYTES))
    for name, definition in SCHEMA.tables.items():
        ws = wb[definition.sheet]
        if not descriptor_row:
            ws.delete_rows(2)
        if not dataset.path(name).exists():
            continue
        header = [cell_text(cell.value) for cell in ws[1]]
        rows = dataset.read(name)
        first = 3 if descriptor_row else 2
        for i, record in enumerate(rows[1:]):
            for column, text in zip(rows[0], record):
                if text:
                    value = excel_value(text, definition.columns[column].type)
                    ws.cell(first + i, header.index(column) + 1, value)
    wb.save(path)
    return path


def edit_workbook(path: Path, change: Callable[[openpyxl.Workbook], None]) -> None:
    wb = openpyxl.load_workbook(path)
    change(wb)
    wb.save(path)


def rename_sheet(path: Path, old: str, new: str) -> None:
    """Renames a sheet of a workbook. Goes through a temporary title, as openpyxl ignores case when comparing titles."""
    def rename(wb: openpyxl.Workbook) -> None:
        wb[old].title = "renaming"
        wb["renaming"].title = new
    edit_workbook(path, rename)


def passes(files: list[tuple[str, Path]]) -> tuple[dict[str, Table], list[str]]:
    """Validates the files, which have to pass. Returns the checked tables and the notices."""
    notices = []
    tables = validate_files(files, SCHEMA, notice=notices.append)
    return tables, notices


def fails(files: list[tuple[str, Path]], message: str) -> None:
    """Checks that validating the files fails with exactly this message."""
    with pytest.raises(ValidationError) as error:
        validate_files(files, SCHEMA)
    assert str(error.value) == message


def as_csv(tables: dict[str, Table]) -> dict[str, str]:
    """Each table as the CSV text the validator writes, to compare tables read from different files."""
    return {name: table.to_csv() for name, table in tables.items()}
