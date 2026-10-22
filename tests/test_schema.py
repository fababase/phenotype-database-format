"""
The schema itself: there is exactly one schema.yaml, mistakes in it are reported, and what is made from it, the
collection sheet and the tables in schema.md, matches it.
"""
import io
import os
import shutil
import subprocess
import sys

import openpyxl
import pytest
import yaml

import schema_md
from helpers import BLANK_SHEET, ROOT, SCHEMA, VALID
from validation import READ_ME_SHEET, SCHEMA_FILE, TYPE_PROBLEMS, cell_text, help_text, load_schema


def test_schema_yaml_has_only_one_copy():
    """schema/schema.yaml is never copied, e.g. into app/, so nothing can be checked against a stale copy."""
    copies = []
    for folder, subfolders, files in os.walk(ROOT):
        subfolders[:] = [name for name in subfolders if name != ".git"]
        copies += [(os.path.relpath(folder, ROOT), name) for name in files
                   if name.lower() in ("schema.yaml", "schema.yml")]
    assert copies == [("schema", "schema.yaml")]
    assert SCHEMA_FILE == ROOT / "schema" / "schema.yaml"


def test_schema_yaml_next_to_the_validator(tmp_path):
    """The browser app has schema.yaml next to validation.py, outside the repository: the deploy Action puts it there."""
    for path in (ROOT / "app" / "validation.py", ROOT / "app" / "requirements.txt", SCHEMA_FILE):
        shutil.copy(path, tmp_path)
    result = subprocess.run([sys.executable, str(tmp_path / "validation.py"), str(VALID)],
                            capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize("change, message", [
    (lambda raw: raw["tables"]["Germplasm"]["columns"]["Name"].update(type="string"),
     "schema.yaml: unknown type 'string' of Germplasm.Name"),
    (lambda raw: raw["tables"]["SeedLot"]["columns"]["GPID"].update(links_to="Germplasms"),
     "schema.yaml: SeedLot.GPID links to unknown table 'Germplasms'"),
    (lambda raw: raw["tables"]["Germplasm"]["columns"]["Name"].pop("description"),
     "schema.yaml: Germplasm.Name has no description"),
    (lambda raw: raw["tables"]["Climate"].pop("needed"),
     "schema.yaml: Climate has no needed"),
    (lambda raw: raw["rules"].append("every_unit_has_a_score"),
     "schema.yaml: unknown rule 'every_unit_has_a_score'"),
], ids=["unknown type", "link to unknown table", "column without a description", "optional table without needed",
        "unknown rule"])
def test_mistake_in_schema_yaml(tmp_path, change, message):
    raw = yaml.safe_load(SCHEMA_FILE.read_text(encoding="utf-8"))
    change(raw)
    path = tmp_path / "schema.yaml"
    path.write_text(yaml.safe_dump(raw, sort_keys=False), encoding="utf-8")
    with pytest.raises(ValueError) as raised:
        load_schema(path)
    assert str(raised.value) == message


# ---------------------------------------------------------------- Collection sheet

def test_collection_sheet_matches_schema():
    """
    The collection sheet has the READ ME 1st sheet, then a sheet for each table, named after it. Each has the
    table's columns in row 1, in the schema's order, and their help text in row 2.
    """
    wb = openpyxl.load_workbook(io.BytesIO(BLANK_SHEET), read_only=True)
    assert wb.sheetnames == [READ_ME_SHEET, *SCHEMA.tables]
    for name, table in SCHEMA.tables.items():
        header, help_row = ([cell_text(value) for value in row]
                            for row in wb[name].iter_rows(max_row=2, values_only=True))
        assert header == list(table.columns), name
        assert help_row == [help_text(column, SCHEMA) for column in table.columns.values()], name
    wb.close()


def test_read_me_sheet_lists_the_version_and_the_tables():
    wb = openpyxl.load_workbook(io.BytesIO(BLANK_SHEET), read_only=True)
    rows = [[cell_text(value) for value in row] for row in wb[READ_ME_SHEET].iter_rows(values_only=True)]
    wb.close()
    assert f"schema version {SCHEMA.version}" in rows[0][0]
    tables = rows[rows.index(["Sheet", "One row is", "Needed", "Filled in, relative to the experiment"]) + 1:]
    assert tables == [[t.name, t.one_row_is, t.needed, t.filled_in] for t in SCHEMA.tables.values()]


@pytest.mark.parametrize("table, column, text", [
    ("Germplasm", "GPID", "Required, unique. Germplasm ID. Only the letters a-z and A-Z, the digits 0-9 and _. "
                          "Example: vandrel"),
    ("Germplasm", "Name", "Optional. Name of the variety or line. Example: Vandrel"),
    ("SeedLot", "GPID", "Required. Germplasm the seed is of. A GPID from Germplasm. Example: vandrel"),
    ("Climate", "UNID", "Optional. Unit the measurement belongs to. Empty or missing means all units in the "
                        "experiment. A UNID from ExperimentalUnit. Example: nordbean_irrigation_northfield_2025_plot101"),
    ("Phenotype", "Score", "Required. The score. A number, with . as the decimal mark. Example: 94"),
    ("Features", "Type", "Optional. Kind of value, in lower case: c for text, i for whole numbers, d for real "
                         "numbers. Empty means c. Example: c"),
    ("PhenotypeDescription", "DescriptionOfMethod", "Required. What is measured, on which part of the plant, at "
                                                    "which growth stage, with what, on how many plants and how they are "
                                                    "chosen, how the score is calculated, and what every code means."),
], ids=["ID", "text", "link", "bold", "number", "code", "no example"])
def test_help_text(table, column, text):
    """
    Whether the column is required, its description without `code` and **bold** marks, a hint for its type, and
    its example if it has one.
    """
    assert help_text(SCHEMA.tables[table].columns[column], SCHEMA) == text


def test_examples_are_valid_values():
    """Each example is text, and a valid value for its column's type, so schema.md and the sheet never show a wrong one."""
    for table in SCHEMA.tables.values():
        for column in table.columns.values():
            if column.example is None:
                continue
            where = f"{table.name}.{column.name}"
            assert isinstance(column.example, str), f"{where}: put the example in quotes in schema.yaml"
            problem = TYPE_PROBLEMS.get(column.type)
            assert problem is None or problem(column.example) is None, f"{where}: {problem(column.example)}"


# ---------------------------------------------------------------- schema.md

def schema_md_text() -> str:
    return (ROOT / "schema" / "schema.md").read_text(encoding="utf-8")


def test_schema_md_has_a_place_for_every_table():
    """
    The Update schema.md Action can write the tables: schema.md has comments to put each one between. Whether they
    are up to date is not checked, as the Action brings them up to date after each push to main.
    """
    schema_md.update(schema_md_text(), SCHEMA)


def test_schema_md_has_the_version():
    assert f"This is **version {SCHEMA.version}** of the schema." in schema_md_text()


def test_schema_md_without_a_place_for_a_table():
    text = schema_md_text().replace("<!-- END columns of Climate -->", "")
    with pytest.raises(ValueError) as raised:
        schema_md.update(text, SCHEMA)
    assert str(raised.value) == ("schema.md has no place for the table 'columns of Climate'. Put "
                                 f"'<!-- BEGIN columns of Climate: {schema_md.NOTE} -->' and "
                                 "'<!-- END columns of Climate -->' where it goes.")
