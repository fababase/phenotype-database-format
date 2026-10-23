"""
Builds the blank collection sheet from the schema: a workbook for filling in the tables by hand. The app offers it
for download, built from the same schema.yaml it checks against, so the two always match. From the command line:

    python collection_sheet.py <file.xlsx>

The workbook has a READ ME 1st sheet, then one sheet per table, named after it, with the column names in row 1 and
what goes in each column in row 2. The validator skips the READ ME 1st sheet, and row 2 when it is there.
"""
import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from validation import FEATURE_TYPES, READ_ME_SHEET, READ_ME_TITLE, Column, Schema, TableDef, help_text, load_schema

# Links in the READ ME 1st sheet. They hold the repository's name, so they change if it is renamed.
SITE = "https://fababase.github.io/phenotype-database-format/"
SCHEMA_DOC = "https://github.com/fababase/phenotype-database-format/blob/main/schema/schema.md"

TITLE = Font(bold=True, size=14)
BOLD = Font(bold=True)
LINK = Font(color="0563C1", underline="single")
HELP = Font(italic=True, color="595959")
REQUIRED = PatternFill("solid", start_color="FFE699")
# Number formats: text, so Excel keeps what is typed, such as 1-2 or 1.5, rather than making it a date; and dates
# shown the way the validator reads them, so a date Excel read the wrong way round is seen at once
TEXT_FORMAT = "@"
DATE_FORMAT = "yyyy-mm-dd"
# The last row of a sheet in Excel. The lists to pick from cover the rows below the help text, down to it.
LAST_ROW = 1048576


def add_read_me(ws, schema: Schema) -> None:
    """
    The READ ME 1st sheet: which schema version the workbook is for, how to fill it in, where to read more, and
    the tables. It is kept short, as schema.md has the rest.
    """
    lines = [
        (READ_ME_TITLE.format(version=schema.version), TITLE, None),
        ("", None, None),
        ("How to fill it in", BOLD, None),
        ("Each sheet is a table. Row 1 holds the column names, and row 2 says what goes in each column. "
         "Start your data in row 3.", None, None),
        ("Do not rename the sheets or the columns. Leave row 2 as it is, or delete it.", None, None),
        ("The names of required columns are coloured. Fill these in on every row.", None, None),
        ("Leave a cell empty when there is no value. Optional columns and tables can be left empty.", None, None),
        ("Cells are formatted as text, so Excel keeps exactly what you type. Dates are shown as YYYY-MM-DD.",
         None, None),
        ("Columns that hold an ID from another sheet list its IDs to pick from.", None, None),
        ("Then check the workbook with the validator, which says what to fix:", None, SITE),
        ("", None, None),
        ("To read more", BOLD, None),
        ("Every table and column, and the rules they follow:", None, SCHEMA_DOC),
        ("How the tables link:", None, f"{SITE}diagram/"),
        ("", None, None),
        ("The tables", BOLD, None),
    ]
    for row, (text, font, link) in enumerate(lines, start=1):
        cell = ws.cell(row, 1, f"{text} {link}" if link else text or None)
        if link:
            cell.hyperlink = link
            font = LINK
        if font:
            cell.font = font

    ws.append(["Sheet", "One row is", "Needed", "Filled in, relative to the experiment"])
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    for table in schema.tables.values():
        ws.append([table.name, table.one_row_is, table.needed, table.filled_in])
    for column, width in zip("ABCD", (24, 60, 32, 36)):
        ws.column_dimensions[column].width = width


def choices(table: TableDef, column: Column) -> str | None:
    """
    The values a column lists to pick from, as the formula of an Excel list: the IDs in the sheet it links to, below
    its help text, or the feature types. None if it lists none.
    """
    if column.links_to:
        return f"'{column.links_to}'!$A$3:$A${LAST_ROW}"
    if (table.name, column.name) == ("Features", "Type"):
        return f'"{",".join(FEATURE_TYPES)}"'
    return None


def add_table(ws, table: TableDef, schema: Schema) -> None:
    """
    A table's sheet: the column names in row 1, with those of required columns coloured, and what goes in each
    column in row 2. Each column is formatted as text, or as a date, and some list the values to pick from.
    """
    columns = list(table.columns.values())
    ws.append([column.name for column in columns])
    ws.append([help_text(column, schema) for column in columns])
    for j, column in enumerate(columns, start=1):
        letter = get_column_letter(j)
        ws.cell(1, j).font = BOLD
        if column.required:
            ws.cell(1, j).fill = REQUIRED
        ws.cell(2, j).font = HELP
        ws.cell(2, j).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[letter].width = 50 if column.type == "long_text" else 25
        ws.column_dimensions[letter].number_format = DATE_FORMAT if column.type == "date" else TEXT_FORMAT
        formula = choices(table, column)
        if formula:
            # Only a help: any value can still be typed, and the validator checks it
            values = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
            values.add(f"{letter}3:{letter}{LAST_ROW}")
            ws.add_data_validation(values)
    # Only row 1 stays in view, as row 2 can be tall
    ws.freeze_panes = "A2"


def build(schema: Schema) -> bytes:
    """The blank collection sheet, as the bytes of an .xlsx file."""
    wb = Workbook()
    wb.active.title = READ_ME_SHEET
    add_read_me(wb.active, schema)
    for table in schema.tables.values():
        add_table(wb.create_sheet(table.name), table, schema)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def main(path: str) -> None:
    Path(path).write_bytes(build(load_schema()))
    print(f"Wrote the blank collection sheet to {path}.")


if __name__ == "__main__":
    main(sys.argv[1])
