"""
Checks phenotype database tables against the schema in schema.yaml.

The tables come either as CSV files, one per table and named after it (Germplasm.csv, SeedLot.csv, ...),
or as one filled-in collection sheet (.xlsx). Checking stops at the first error, with a message that says
where it is and how to fix it. Warnings point out what is allowed, but more often a mistake, such as a very early
date; the tables pass with them.

From the command line:
    python validation.py <CSV files, a folder of CSV files, or a .xlsx file> [--out FOLDER]
"""
import argparse
import csv
import difflib
import io
import math
import re
import string
import sys
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from importlib import metadata
from pathlib import Path

import openpyxl
import pandas as pd
import yaml

HERE = Path(__file__).resolve().parent
# The only copy of the schema. It is never copied anywhere else in the repository, so nothing can be checked
# against a stale one. Only the browser app has it next to this file, where the deploy Action puts it.
SCHEMA_FILE = HERE.parent / "schema" / "schema.yaml"
if not SCHEMA_FILE.exists():
    SCHEMA_FILE = HERE / "schema.yaml"
# The package versions that the tests and the browser app run with
REQUIREMENTS_FILE = HERE / "requirements.txt"

COLUMN_TYPES = {"id", "text", "long_text", "list", "integer", "number", "date", "coordinates", "features"}
# Feature types, in lower case: c = text, i = whole number, d = real number. An empty Type means c.
FEATURE_TYPES = ("c", "i", "d")

# Digits are only 0-9, so that full-width or Arabic-Indic digits, which Python reads as numbers too, do not pass
DECIMAL = r"-?[0-9]+(?:\.[0-9]+)?"
INTEGER = re.compile(r"-?[0-9]+")
NUMBER = re.compile(rf"{DECIMAL}(?:[eE][-+]?[0-9]+)?")
DATE = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}")
COORDINATES = re.compile(rf"{DECIMAL};{DECIMAL}")
# min..max, where an end left blank means no limit
RANGE = re.compile(rf"({NUMBER.pattern})?\.\.({NUMBER.pattern})?")
# A number with one ',' that could be a decimal mark or separate thousands, such as 1,234
THOUSANDS = re.compile(r"-?[1-9][0-9]{0,2},[0-9]{3}")
ID = re.compile(r"[A-Za-z][A-Za-z0-9_]*")
ID_CHARACTER = re.compile(r"[A-Za-z0-9_]")
# Besides letters, the digits 0-9 and spaces, text can hold the symbols of a US keyboard (ASCII), apart from ';',
# which separates the values of a list, and '=', which joins an FTID to its value in Features cells. Other symbols,
# such as '°', are garbled when a file is read in the wrong encoding.
TEXT_SYMBOLS = set(string.punctuation) - {";", "="}
# Dates before this get a warning, as they are more often a mistake, such as 0025 for 2025, than a real date
EARLY = date(1950, 1, 1)

# CSV separators to recognise, and how to name them in notices
DELIMITERS = {",": "commas", ";": "semicolons", "\t": "tabs"}


class ValidationError(ValueError):
    """
    Raised when the data breaks a rule of the schema.
    Any other exception from the validation is a bug in the validator, not a problem with the data.
    """


# ---------------------------------------------------------------- Schema

@dataclass
class Column:
    name: str
    type: str
    required: bool = False
    unique: bool = False
    links_to: str | None = None
    description: str = ""      # may have `code` and **bold** in it
    example: str | None = None


@dataclass
class TableDef:
    name: str
    code: str
    required: bool
    columns: dict[str, Column]
    one_row_is: str  # what one row of the table is
    filled_in: str   # when the table is filled in, relative to the experiment
    needed: str      # when the table is needed: "always" if it is required

    @property
    def id_column(self) -> str:
        """The ID column, which is always the first column."""
        return next(iter(self.columns))

    def add_hint(self) -> str:
        return f"Add {self.name}.csv, or the sheet '{self.name}' to the workbook."


@dataclass
class Schema:
    version: str
    missing_values: set[str]  # lower case, including "" for an empty cell
    tables: dict[str, TableDef]
    rules: list[str]

    def match(self, name: str) -> str | None:
        """The table that a file stem or sheet name stands for: its table name or code, in any case."""
        key = name.strip().lower()
        for table in self.tables.values():
            if key in (table.name.lower(), table.code.lower()):
                return table.name
        return None


def load_schema(path: Path = SCHEMA_FILE) -> Schema:
    """Reads the schema from schema.yaml. Raises ValueError if the file itself is inconsistent."""
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    tables = {}
    for name, table in raw["tables"].items():
        columns = {column: Column(column, **spec) for column, spec in table["columns"].items()}
        required = table.get("required", False)
        tables[name] = TableDef(name, table["code"], required, columns, table.get("one_row_is", ""),
                                table.get("filled_in", ""), "always" if required else table.get("needed", ""))
    for table in tables.values():
        for key in ("one_row_is", "filled_in", "needed"):
            if not getattr(table, key):
                raise ValueError(f"schema.yaml: {table.name} has no {key}")
        for column in table.columns.values():
            if column.type not in COLUMN_TYPES:
                raise ValueError(f"schema.yaml: unknown type '{column.type}' of {table.name}.{column.name}")
            if column.links_to and column.links_to not in tables:
                raise ValueError(f"schema.yaml: {table.name}.{column.name} links to unknown table '{column.links_to}'")
            if not column.description:
                raise ValueError(f"schema.yaml: {table.name}.{column.name} has no description")
    for rule in raw.get("rules", []):
        if rule not in RULES:
            raise ValueError(f"schema.yaml: unknown rule '{rule}'")
    return Schema(
        version=str(raw["version"]),
        missing_values={""} | {str(value).strip().lower() for value in raw["missing_values"]},
        tables=tables,
        rules=raw.get("rules", []),
    )


# ---------------------------------------------------------------- Collection sheet

# The collection sheet's sheet of instructions, which the validator skips. Compared ignoring case.
READ_ME_SHEET = "READ ME 1st"
# The first cell of the READ ME 1st sheet, which the validator reads the version from
READ_ME_TITLE = "Collection sheet for the FabaBase phenotype database, schema version {version}"
SHEET_VERSION = re.compile(r"schema version (\S+)")

# What row 2 of the collection sheet says about each type of column, after the column's description
TYPE_HINTS = {
    "id": "Only the letters a-z and A-Z, the digits 0-9 and _, starting with a letter.",
    "text": "Letters, the digits 0-9, spaces and the symbols of a US keyboard except ; and =.",
    "long_text": "Letters, the digits 0-9, spaces, line breaks and the symbols of a US keyboard except ; and =.",
    "list": "Values separated by ;, with no spaces around it.",
    "integer": "A whole number.",
    "number": "A number, with . as the decimal mark.",
    "date": "A date, written YYYY-MM-DD, or a date that Excel recognises as one.",
    "coordinates": "latitude;longitude in decimal degrees, without spaces.",
    "features": "FTID1=value1;FTID2=value2, with FTIDs from Features, and no spaces around = or ;.",
}


def plain(text: str) -> str:
    """A description without its `code` and **bold** marks, as the collection sheet shows it."""
    return text.replace("**", "").replace("`", "")


def help_start(column: Column) -> str:
    """How row 2 of the collection sheet starts for a column: whether it is required, then its description."""
    need = ("Required" if column.required else "Optional") + (", unique" if column.unique else "")
    return f"{need}. {plain(column.description)}"


def help_text(column: Column, schema: Schema) -> str:
    """What row 2 of the collection sheet says about a column."""
    if column.links_to:
        hint = f"A {schema.tables[column.links_to].id_column} from {column.links_to}."
    else:
        hint = TYPE_HINTS.get(column.type, "")
    # Without a full stop after it, which could be taken as part of the value
    example = f"Example: {column.example}" if column.example else ""
    return " ".join(part for part in (help_start(column), hint, example) if part)


# ---------------------------------------------------------------- Tables

@dataclass
class RawTable:
    """A table as read from a file or sheet, before any checks."""
    source: str             # "Germplasm.csv" or "sheet 'Germplasm'", for messages
    header: list[str]
    records: list[list[str]]
    row_numbers: list[int]  # row number of each record in the file or sheet; the header is row 1


@dataclass
class Table:
    """A table whose columns match the schema. All values are text, in the schema's column order."""
    definition: TableDef
    source: str
    data: pd.DataFrame
    missing: pd.DataFrame  # True where a cell counts as missing
    row_numbers: list[int]

    @property
    def name(self) -> str:
        return self.definition.name

    def where(self, i: int, column: str) -> str:
        return f"{self.source}, row {self.row_numbers[i]}, column {column}"

    def values(self, column: str) -> pd.Series:
        """The values of a column that are not missing."""
        return self.data[column][~self.missing[column]]

    def dates(self, column: str) -> list[date | None]:
        """The dates in a date column whose values have been checked, with None where a value is missing."""
        return [None if m else date.fromisoformat(v) for v, m in zip(self.data[column], self.missing[column])]

    def to_csv(self) -> str:
        """The table as CSV: comma-separated, a header row, dates YYYY-MM-DD and missing values as empty cells."""
        return self.data.mask(self.missing, "").to_csv(index=False, lineterminator="\n")


def cell_text(value) -> str:
    """The text of an Excel cell. Dates become YYYY-MM-DD, and numbers are written without exponents."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(repr(value)), "f")
    return str(value)


def read_csv(name: str, path: Path, notice: Callable[[str], None]) -> RawTable:
    """
    Reads a CSV file in UTF-8 or Windows encoding, separated by commas, semicolons or tabs.
    Every row has as many cells as the header, apart from blank lines.
    """
    content = path.read_bytes()
    if b"\0" in content:
        raise ValidationError(f"{name} contains null characters, so it is not a plain CSV file. "
                              "Save it as 'CSV UTF-8'.")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1252")
        except UnicodeDecodeError:
            raise ValidationError(f"{name}: could not read the text of the file. Save it as 'CSV UTF-8'.") from None
        notice(f"{name} is not UTF-8, so it was read as Windows text.")
    first_line = text.split("\n", 1)[0]
    delimiter = max(DELIMITERS, key=first_line.count)
    if delimiter != ",":
        notice(f"{name} is separated by {DELIMITERS[delimiter]}.")
    # Strict, so that a quote mark that is never closed is an error, rather than swallowing the rows after it
    rows = []
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True):
            rows.append(row)
    except csv.Error as e:
        raise ValidationError(f"{name}, row {len(rows) + 1}: could not read the row as CSV ({e}). Check its quote "
                              "marks: a value that starts with \" has to end with \" too.") from None
    if not rows:
        raise ValidationError(f"{name} is empty. Its first row has to hold the column names.")
    width = len(rows[0])
    separator = DELIMITERS[delimiter]
    for number, row in enumerate(rows[1:], start=2):
        if row and len(row) != width:
            hint = (f"Put quotes around any value with {separator} in it, or add the missing column names to row 1."
                    if len(row) > width else f"Each row needs a cell for every column, with {separator} between "
                    "the cells even when they are empty.")
            raise ValidationError(f"{name}, row {number} has {len(row)} cells, but row 1 has {width}. {hint}")
    return RawTable(name, rows[0], rows[1:], list(range(2, len(rows) + 1)))


def check_version(name: str, ws, schema: Schema) -> None:
    """
    A collection sheet has to be for the schema version it is checked against. Its READ ME 1st sheet says which
    version it is for; if it does not, nothing is checked.
    """
    first_row = next(ws.iter_rows(max_row=1, max_col=1, values_only=True), ())
    match = SHEET_VERSION.search(cell_text(first_row[0]) if first_row else "")
    if match and match[1] != schema.version:
        raise ValidationError(f"{name} is a collection sheet for schema version {match[1]}, but it is checked "
                              f"against schema version {schema.version}. Download a new blank collection sheet, "
                              "and copy your data into it.")


def read_workbook(name: str, path: Path, schema: Schema, notice: Callable[[str], None],
                  warning: Callable[[str], None]) -> list[tuple[str, RawTable]]:
    """
    Reads the sheets of a collection sheet. Returns (table name, table) for each sheet that is a table.
    A hidden sheet that is a table is read as usual, with a warning, and any other hidden sheet is skipped.
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise ValidationError(f"Could not open {name} as an Excel workbook ({e}).") from None
    try:
        read_me = [ws for ws in wb.worksheets if ws.title.strip().lower() == READ_ME_SHEET.lower()]
        if read_me:
            check_version(name, read_me[0], schema)
        found = []
        for ws in wb.worksheets:
            if ws in read_me:
                continue
            source = f"sheet '{ws.title}'"
            table = schema.match(ws.title)
            hidden = ws.sheet_state != "visible"
            if table is None:
                if hidden:
                    notice(f"{source} is hidden and is not a table, so it was skipped.")
                    continue
                raise ValidationError(f"{source} is not a table of the schema. Rename it after its table "
                                      f"({', '.join(schema.tables)}), or delete it.")
            if hidden:
                warning(f"{source} is hidden, but it is the {table} table, so it was checked.")
            rows = [[cell_text(value) for value in row] for row in ws.iter_rows(values_only=True)]
            if not rows:
                raise ValidationError(f"{source} is empty. Put the column names in row 1, or delete the sheet.")
            # Skip row 2 if it holds the collection sheet's help text, which no ID can start with. Only the start
            # is compared, so blank sheets from earlier versions pass even if a type hint or example was added since.
            definition = schema.tables[table]
            start_of_help = help_start(definition.columns[definition.id_column]).lower()
            start = 1
            if len(rows) > 1 and rows[1] and rows[1][0].strip().lower().startswith(start_of_help):
                start = 2
            found.append((table, RawTable(source, rows[0], rows[start:], list(range(start + 1, len(rows) + 1)))))
        return found
    finally:
        wb.close()


def load_files(files: list[tuple[str, Path]], schema: Schema, notice: Callable[[str], None],
               warning: Callable[[str], None]) -> dict[str, RawTable]:
    """
    Reads the tables from CSV files or one collection sheet.
    Parameters:
        files: (file name, path) of each file. The file name says which table a CSV file is.
    Returns:
        dict: table name -> table as read.
    """
    if not files:
        raise ValidationError("No files given.")
    for name, _ in files:
        if Path(name).suffix.lower() not in (".csv", ".xlsx"):
            raise ValidationError(f"{name} is not a CSV or .xlsx file. "
                                  "Give one CSV file per table, or one collection sheet (.xlsx).")
    workbooks = [file for file in files if Path(file[0]).suffix.lower() == ".xlsx"]
    if workbooks and len(files) > 1:
        raise ValidationError("Give either one collection sheet (.xlsx) or CSV files, not both or several workbooks.")

    if workbooks:
        found = read_workbook(*workbooks[0], schema, notice, warning)
    else:
        found = []
        for name, path in files:
            table = schema.match(Path(name).stem)
            if table is None:
                raise ValidationError(f"{name} is not named after a table of the schema. Name each file after "
                                      f"its table: {', '.join(f'{t}.csv' for t in schema.tables)}.")
            found.append((table, read_csv(name, path, notice)))

    raws = {}
    for table, raw in found:
        if table in raws:
            raise ValidationError(f"{raw.source} and {raws[table].source} are both the {table} table. Keep only one.")
        raws[table] = raw
    return raws


def build_table(definition: TableDef, raw: RawTable, schema: Schema) -> Table:
    """
    Matches the columns of a table to the schema: in any order, optional columns may be left out,
    and unknown columns are an error. A row in which every cell is missing has no values: at the end of the table
    it is dropped, as Excel often leaves such rows there unseen, but with rows after it, it is an error, as it may
    end the table, with something else below it.
    """
    width = max([len(raw.header)] + [len(record) for record in raw.records])
    header = [name.strip() for name in raw.header] + [""] * (width - len(raw.header))
    records = [record + [""] * (width - len(record)) for record in raw.records]
    blank = [all(value.strip().lower() in schema.missing_values for value in record) for record in records]
    end = len(blank)
    while end and blank[end - 1]:
        end -= 1
    records = records[:end]
    row_numbers = raw.row_numbers[:end]

    positions = {}
    for j, name in enumerate(header):
        if not name:
            if any(record[j].strip() for record in records):
                raise ValidationError(f"{raw.source}, column {j + 1} has values but no column name in row 1. "
                                      "Add the column name, or delete the column.")
            continue
        if name not in definition.columns:
            # Compared ignoring case, so that gpid gets GPID
            names = {column.lower(): column for column in definition.columns}
            close = difflib.get_close_matches(name.lower(), names, n=1)
            hint = (f"Did you mean '{names[close[0]]}'?" if close
                    else f"The columns of {definition.name} are: {', '.join(definition.columns)}.")
            raise ValidationError(f"{raw.source}: '{name}' is not a column of {definition.name}. {hint}")
        if name in positions:
            raise ValidationError(f"{raw.source}: the column '{name}' appears twice.")
        positions[name] = j
    for column in definition.columns.values():
        if column.required and column.name not in positions:
            raise ValidationError(f"{raw.source}: the column '{column.name}' is missing. "
                                  f"It is required in {definition.name}.")
    if any(blank[:end]):
        raise ValidationError(f"{raw.source}, row {raw.row_numbers[blank.index(True)]} has no values, but rows after "
                              "it do. Delete the row. If the rows below it are not part of the table, delete them too.")

    data = pd.DataFrame({name: [record[positions[name]] if name in positions else "" for record in records]
                         for name in definition.columns}, dtype=str)
    missing = pd.DataFrame({name: data[name].str.strip().str.lower().isin(schema.missing_values)
                            for name in definition.columns})
    return Table(definition, raw.source, data, missing, row_numbers)


# ---------------------------------------------------------------- Checks

def first_row(table: Table, column: str, mask: pd.Series, describe: Callable[[int], str]) -> str | None:
    """
    A message about the first row where mask is True, which also says how many more rows are like it, or None if
    there is no such row. describe(i) says what is wrong in row i.
    """
    rows = mask.to_numpy().nonzero()[0]
    if len(rows) == 0:
        return None
    i = int(rows[0])
    others = len(rows) - 1
    more = ""
    if others == 1:
        more = " (1 more row in this column has the same problem.)"
    elif others > 1:
        more = f" ({others} more rows in this column have the same problem.)"
    return f"{table.where(i, column)}: {describe(i)}{more}"


def fail_first(table: Table, column: str, mask: pd.Series, describe: Callable[[int], str]) -> None:
    """Raises a ValidationError for the first row where mask is True. describe(i) says what is wrong in row i."""
    message = first_row(table, column, mask, describe)
    if message:
        raise ValidationError(message)


def warn_first(table: Table, column: str, mask: pd.Series, describe: Callable[[int], str],
               warning: Callable[[str], None]) -> None:
    """Gives one warning for the rows where mask is True, about the first of them."""
    message = first_row(table, column, mask, describe)
    if message:
        warning(message)


def character_name(c: str) -> str:
    """How messages name a character: in quotes, or in words if it cannot be seen."""
    if c == " ":
        return "a space"
    if c in "\n\r":
        return "a line break"
    if c == "\t":
        return "a tab"
    if c.isprintable() and not c.isspace():
        return f"'{c}'"
    name = unicodedata.name(c, "")
    return f"the invisible character U+{ord(c):04X}" + (f" ({name.lower()})" if name else "")


def character_names(characters: list[str]) -> str:
    """Characters named in a list, such as "'-', '/' and a space"."""
    names = [character_name(c) for c in characters]
    return names[0] if len(names) == 1 else f"{', '.join(names[:-1])} and {names[-1]}"


def id_problem(value: str) -> str | None:
    if ID.fullmatch(value):
        return None
    if ID.fullmatch(value.strip()):
        return f"'{value}' has spaces before or after it. Remove them."
    bad = list(dict.fromkeys(c for c in value.strip() if not ID_CHARACTER.fullmatch(c)))
    if not bad:
        return f"'{value}' starts with '{value.strip()[0]}'. IDs have to start with a letter, a-z or A-Z."
    message = (f"'{value}' contains {character_names(bad)}. IDs can only hold the letters a-z and A-Z without "
               "accents, the digits 0-9 and '_'.")
    # Letters such as 'ø' have no obvious replacement, so a fix is only suggested for other characters, and only
    # if it is a valid ID
    fix = re.sub(r"[^A-Za-z0-9_]+", "_", value.strip())
    if any(c.isalpha() for c in bad) or not ID.fullmatch(fix):
        return message
    return f"{message} Use '_' instead: {fix}"


def other_digits(value: str) -> str | None:
    """What is wrong with a value that has digits other than 0-9, such as full-width or Arabic-Indic ones."""
    digits = [c for c in dict.fromkeys(value) if c.isdecimal() and not c.isascii()]
    if digits:
        return f"'{value}' has digits other than 0-9: {character_names(digits)}. Write them as 0-9."
    return None


def too_large(value: str) -> str | None:
    """What is wrong with a number so large that it is read as infinity, such as 1e999."""
    if math.isinf(float(value)):
        return f"'{value}' is too large a number: it would be read as infinity."
    return None


def swap_hint(day: date, fits: Callable[[date], bool]) -> str:
    """
    For a message about a date that breaks a rule: what the date would be with day and month swapped, if that is a
    real date that is not in the future and fits(it). An empty text otherwise.
    """
    try:
        other = date(day.year, day.day, day.month)
    except ValueError:
        return ""
    if other == day or other > date.today() or not fits(other):
        return ""
    return f" With day and month swapped, it would be {other.isoformat()}, which fits. Check which is meant."


def integer_problem(value: str) -> str | None:
    if INTEGER.fullmatch(value):
        return None
    return other_digits(value) or f"'{value}' is not a whole number. Write digits only, e.g. 12."


def number_problem(value: str) -> str | None:
    """A fix is only suggested when it is the one way to read the value."""
    if NUMBER.fullmatch(value):
        return too_large(value)
    digits = other_digits(value)
    if digits:
        return digits
    if THOUSANDS.fullmatch(value):
        return (f"'{value}' has ','. Numbers have '.' as the decimal mark and no separator between thousands. "
                f"Write {value.replace(',', '.')} if ',' is the decimal mark, or {value.replace(',', '')} if it "
                "separates thousands. These are different numbers, so check which one is meant.")
    if NUMBER.fullmatch(value.replace(",", ".")):
        return (f"'{value}' has ','. Numbers have '.' as the decimal mark and no thousands separator, "
                f"so write {value.replace(',', '.')}")
    return (f"'{value}' is not a number. Write only the number, with '.' as the decimal mark and no thousands "
            "separator, units or words, e.g. 3.5.")


def date_problem(value: str) -> str | None:
    if DATE.fullmatch(value):
        try:
            day = date.fromisoformat(value)
        except ValueError:
            return f"'{value}' is not a real date."
        if day > date.today():
            return f"'{value}' is in the future.{swap_hint(day, lambda other: True)}"
        return None
    if DATE.match(value):
        return f"'{value}' holds more than a date. Write only the date, as YYYY-MM-DD."
    return other_digits(value) or f"'{value}' is not a date written YYYY-MM-DD, e.g. 2023-05-17."


def coordinates_problem(value: str) -> str | None:
    if COORDINATES.fullmatch(value):
        latitude, longitude = (float(part) for part in value.split(";"))
        if -90 <= latitude <= 90 and -180 <= longitude <= 180:
            return None
        return (f"'{value}' is outside the range of coordinates: latitude goes from -90 to 90, and longitude from "
                "-180 to 180.")
    digits = other_digits(value)
    if digits:
        return digits
    hint = " Remove the spaces." if COORDINATES.fullmatch(re.sub(r"\s", "", value)) else ""
    return f"'{value}' is not written latitude;longitude in decimal degrees, e.g. 56.170308;10.207013.{hint}"


def text_problem(value: str, kind: str = "text", subject: str = "the text") -> str | None:
    """
    What is wrong with text, or None if nothing is. Text holds only letters, the digits 0-9, spaces and
    TEXT_SYMBOLS. Long text can hold line breaks too, and a list ';' between its values.
    subject is how the message starts: "the text", or the value in quotes.
    """
    allowed = TEXT_SYMBOLS | {"text": {" "}, "long_text": {" ", "\n", "\r"}, "list": {" ", ";"}}[kind]
    bad = [c for c in dict.fromkeys(value) if not (c.isalpha() or c in "0123456789" or c in allowed)]
    if not bad:
        return None
    breaks = ", line breaks" if kind == "long_text" else ""
    return (f"{subject} contains {character_names(bad)}. Text can only hold letters, the digits 0-9, spaces{breaks} "
            "and the symbols of a US keyboard, apart from ';' and '=', which have meanings of their own. Write "
            "other symbols in words, such as 'degrees' for '°'.")


def list_problem(value: str) -> str | None:
    """A list holds values separated by ';', with no spaces around it, and each value is text."""
    if any(item == "" or item != item.strip() for item in value.split(";")):
        return (f"'{value}' is not a list written value1;value2. Separate the values with ';', with no spaces "
                "around it and no empty values.")
    return text_problem(value, "list")


# What is wrong with a value of a type, or None if nothing is. Features cells are checked separately.
TYPE_PROBLEMS: dict[str, Callable[[str], str | None]] = {
    "id": id_problem,
    "text": text_problem,
    "long_text": lambda value: text_problem(value, "long_text"),
    "list": list_problem,
    "integer": integer_problem,
    "number": number_problem,
    "date": date_problem,
    "coordinates": coordinates_problem,
}


def check_cells(table: Table, warning: Callable[[str], None]) -> None:
    """
    Checks each column: required values, types and unique values. IDs are unique ignoring case.
    Dates before EARLY get a warning.
    """
    for column in table.definition.columns.values():
        values = table.data[column.name]
        missing = table.missing[column.name]
        if column.required:
            fail_first(table, column.name, missing, lambda i: (
                f"'{values.iat[i]}' counts as a missing value" if values.iat[i].strip() else "the cell is empty"
            ) + f", but {column.name} is required. Fill it in.")
        problem = TYPE_PROBLEMS.get(column.type)
        if problem:
            problems = pd.Series([None if m else problem(v) for v, m in zip(values, missing)], dtype=object)
            fail_first(table, column.name, problems.notna(), lambda i: problems.iat[i])
        if column.type == "date":
            early = pd.Series([day is not None and day < EARLY for day in table.dates(column.name)], dtype=bool)
            warn_first(table, column.name, early, lambda i: (
                f"'{values.iat[i]}' is before {EARLY.year}. Check that it is right."), warning)
        if column.unique:
            keys = values.str.lower() if column.type == "id" else values

            def describe(i: int) -> str:
                first = int((keys == keys.iat[i]).to_numpy().nonzero()[0][0])
                where = f"'{values.iat[i]}' is also in row {table.row_numbers[first]}"
                if values.iat[first] == values.iat[i]:
                    return f"{where}. Each {column.name} has to be unique."
                return (f"{where}, written '{values.iat[first]}'. Each {column.name} has to be unique, "
                        "and upper and lower case do not count as different.")
            fail_first(table, column.name, keys.duplicated() & ~missing, describe)


def check_links(table: Table, column: Column, tables: dict[str, Table], schema: Schema) -> None:
    """Checks that each value of a column that refers to another table is an ID in that table."""
    values = table.data[column.name]
    filled = ~table.missing[column.name]
    target = tables.get(column.links_to)
    if target is None:
        fail_first(table, column.name, filled, lambda i: (
            f"'{values.iat[i]}' refers to {column.links_to}, but there is no {column.links_to} table. "
            + schema.tables[column.links_to].add_hint()))
        return
    id_column = target.definition.id_column
    ids = set(target.values(id_column))
    loose = {i.strip().lower(): i for i in ids}

    def describe(i: int) -> str:
        value = values.iat[i]
        near = loose.get(value.strip().lower())
        hint = f"Did you mean '{near}'?" if near else f"Check the spelling, or add it to {target.name}."
        return (f"'{value}' is not in column {id_column} of {target.source}. "
                f"IDs have to match exactly, character for character. {hint}")

    fail_first(table, column.name, filled & ~values.isin(ids), describe)


def check_feature_definitions(table: Table | None) -> dict[str, tuple[str, str | None]] | None:
    """
    Checks the Features table: Type has to be c, i, d or empty, and Vocabulary, which `check_cells` has checked as
    a list, has to be a min..max range for i and d. FTIDs are checked as IDs by `check_cells`.
    Returns:
        dict: FTID -> (Type, Vocabulary), with an empty Type given as "c" and an empty Vocabulary as None.
        None if there is no Features table, and an empty dict if it has no rows.
    """
    if table is None:
        return None
    features = {}
    for i in range(len(table.data)):
        row = table.data.iloc[i]
        missing = table.missing.iloc[i]
        ftid = row["FTID"]
        ftype = "c" if missing["Type"] else row["Type"]
        if ftype not in FEATURE_TYPES:
            hint = (f" Types are written in lower case: {ftype.lower()}." if ftype.lower() in FEATURE_TYPES
                    else " Use c (text), i (whole number) or d (real number), or leave it empty for text.")
            raise ValidationError(f"{table.where(i, 'Type')}: '{ftype}' is not a feature type.{hint}")

        vocabulary = None if missing["Vocabulary"] else row["Vocabulary"]
        if vocabulary is not None and ftype != "c":
            ends = vocabulary_range(vocabulary)
            if ends is None or (None not in ends and ends[0] > ends[1]):
                raise ValidationError(f"{table.where(i, 'Vocabulary')}: '{vocabulary}' is not a valid "
                                      f"vocabulary. For Type {ftype}, give a range min..max with min no larger "
                                      "than max, e.g. 0..100. Leave an end blank for no limit, e.g. 0..")

        features[ftid] = (ftype, vocabulary)
    return features


def vocabulary_range(vocabulary: str) -> tuple[float | None, float | None] | None:
    """
    The ends of a range min..max, with None for an end left blank. None if it is not a range, or an end is so large
    that it is read as infinity.
    """
    match = RANGE.fullmatch(vocabulary)
    if not match:
        return None
    ends = tuple(None if end is None else float(end) for end in match.groups())
    if any(end is not None and math.isinf(end) for end in ends):
        return None
    return ends


def check_feature_cells(table: Table, column: str, features: dict[str, tuple[str, str | None]] | None,
                        schema: Schema) -> None:
    """
    Checks a Features column. A cell is empty, or holds entries FTID1=value1;FTID2=value2 with no spaces
    around '=' or ';'. Each FTID has to be in the Features table and can appear only once per cell,
    and each value has to fit its feature's Type and Vocabulary.
    features is what `check_feature_definitions` returned: None if there is no Features table.
    """
    for i, (cell, missing) in enumerate(zip(table.data[column], table.missing[column])):
        if missing:
            continue
        where = f"{table.where(i, column)}: invalid Features cell '{cell}'"
        seen = set()
        for entry in cell.split(";"):
            ftid, sep, value = entry.partition("=")
            if not sep or not ftid or not value or "=" in value or ftid != ftid.strip() or value != value.strip():
                raise ValidationError(f"{where}. Write each entry as FTID=value, separated by ';', "
                                      "without spaces around '=' or ';'.")
            if ftid not in (features or {}):
                hint = f" There is no Features table. {schema.tables['Features'].add_hint()}" if features is None else ""
                raise ValidationError(f"{where}. '{ftid}' is not an FTID in the Features table.{hint}")
            if ftid in seen:
                raise ValidationError(f"{where}. '{ftid}' appears more than once.")
            seen.add(ftid)
            problem = feature_value_problem(value, ftid, *features[ftid])
            if problem:
                raise ValidationError(f"{where}. {problem}")


def feature_value_problem(value: str, ftid: str, ftype: str, vocabulary: str | None) -> str | None:
    """
    Returns what is wrong with a value of a feature, or None if it fits the feature's Type and Vocabulary.
    The Vocabulary is assumed to have passed `check_feature_definitions`.
    """
    if ftype == "i" and not INTEGER.fullmatch(value):
        return other_digits(value) or f"'{value}' is not a whole number, as feature {ftid} (Type i) needs."
    if ftype == "d":
        if not NUMBER.fullmatch(value):
            return other_digits(value) or (f"'{value}' is not a number with '.' as the decimal mark, as feature "
                                           f"{ftid} (Type d) needs.")
        problem = too_large(value)
        if problem:
            return problem
    if ftype == "c":
        problem = text_problem(value, subject=f"'{value}'")
        if problem:
            return problem
    if vocabulary is None:
        return None
    if ftype == "c":
        if value not in vocabulary.split(";"):
            return f"'{value}' is not in the vocabulary of feature {ftid} ({vocabulary})."
    else:
        low, high = vocabulary_range(vocabulary)
        if (low is not None and float(value) < low) or (high is not None and float(value) > high):
            return f"{value} is outside the range of feature {ftid} ({vocabulary})."
    return None


# ---------------------------------------------------------------- Rules that span tables

def climate_unit_in_experiment(tables: dict[str, Table]) -> None:
    """A Climate UNID that is filled in has to belong to the row's EXPID."""
    climate = tables.get("Climate")
    if climate is None:
        return
    units = tables["ExperimentalUnit"]
    unit_experiment = dict(zip(units.data["UNID"], units.data["EXPID"]))
    unids = climate.data["UNID"]
    expids = climate.data["EXPID"]
    wrong = ~climate.missing["UNID"] & (unids.map(unit_experiment) != expids)
    fail_first(climate, "UNID", wrong, lambda i: (
        f"unit '{unids.iat[i]}' belongs to experiment '{unit_experiment[unids.iat[i]]}', but this row's EXPID "
        f"is '{expids.iat[i]}'. Fix the EXPID or the UNID, or leave UNID empty if the value applies to all "
        "units in the experiment."))


def one_intervention_per_unit(tables: dict[str, Table]) -> None:
    """
    In an experiment with interventions, every unit has a row in ExperimentIntervention.
    That it has no more than one is checked by UNID being unique there.
    """
    interventions = tables.get("ExperimentIntervention")
    if interventions is None:
        return
    units = tables["ExperimentalUnit"]
    unit_experiment = dict(zip(units.data["UNID"], units.data["EXPID"]))
    with_row = set(interventions.values("UNID"))
    experiments = {unit_experiment[unid] for unid in with_row}
    lacking = units.data["EXPID"].isin(experiments) & ~units.data["UNID"].isin(with_row)
    fail_first(units, "UNID", lacking, lambda i: (
        f"unit '{units.data['UNID'].iat[i]}' has no row in {interventions.source}, but other units in "
        f"experiment '{units.data['EXPID'].iat[i]}' do. In experiments with interventions, every unit needs "
        "one row, control units included."))


def every_experiment_has_a_unit(tables: dict[str, Table]) -> None:
    """Every experiment has at least one unit."""
    experiments = tables["Experiment"]
    units = tables["ExperimentalUnit"]
    expids = experiments.data["EXPID"]
    without = ~experiments.missing["EXPID"] & ~expids.isin(set(units.values("EXPID")))
    fail_first(experiments, "EXPID", without, lambda i: (
        f"experiment '{expids.iat[i]}' has no units in {units.source}. Every experiment needs at least one unit."))


def no_seed_lot_descends_from_itself(tables: dict[str, Table]) -> None:
    """Going from seed lot to parent seed lot, by ParentSLID, never leads back to the seed lot started from."""
    seed_lots = tables["SeedLot"]
    slids = seed_lots.data["SLID"]
    parents = {slid: parent for slid, parent, missing
               in zip(slids, seed_lots.data["ParentSLID"], seed_lots.missing["ParentSLID"]) if not missing}

    def loop(slid: str) -> list[str] | None:
        """The seed lots from slid back to itself, or None if its ancestors do not lead back to it."""
        path = [slid]
        while path[-1] in parents:
            parent = parents[path[-1]]
            if parent == slid:
                return path + [slid]
            if parent in path:
                return None
            path.append(parent)
        return None

    loops = [loop(slid) for slid in slids]
    fail_first(seed_lots, "ParentSLID", pd.Series([path is not None for path in loops], dtype=bool), lambda i: (
        f"seed lot '{slids.iat[i]}' descends from itself: {' -> '.join(loops[i])}. A seed lot cannot be grown from "
        "itself or its own offspring. Fix the ParentSLID of one of these seed lots."))


def parent_seed_lot_harvested_first(tables: dict[str, Table]) -> None:
    """A seed lot is harvested after the seed lot it was grown from."""
    seed_lots = tables["SeedLot"]
    days = seed_lots.dates("HarvestDate")
    harvested = dict(zip(seed_lots.data["SLID"], days))
    parents = [None if missing else parent
               for parent, missing in zip(seed_lots.data["ParentSLID"], seed_lots.missing["ParentSLID"])]
    parent_days = [harvested[parent] if parent else None for parent in parents]
    wrong = pd.Series([day is not None and parent_day is not None and day <= parent_day
                       for day, parent_day in zip(days, parent_days)], dtype=bool)
    fail_first(seed_lots, "HarvestDate", wrong, lambda i: (
        f"'{days[i]}' is not after {parent_days[i]}, the HarvestDate of its parent seed lot '{parents[i]}'. "
        "A seed lot is harvested after the seed lot it was grown from."
        + swap_hint(days[i], lambda day: day > parent_days[i])))


def experiment_starts_before_it_ends(tables: dict[str, Table]) -> None:
    """An experiment's StartDate is not after its EndDate."""
    experiments = tables["Experiment"]
    starts = experiments.dates("StartDate")
    ends = experiments.dates("EndDate")
    wrong = pd.Series([start is not None and end is not None and start > end for start, end in zip(starts, ends)],
                      dtype=bool)
    fail_first(experiments, "StartDate", wrong, lambda i: (
        f"'{starts[i]}' is after the EndDate, {ends[i]}." + swap_hint(starts[i], lambda day: day <= ends[i])))


def date_within_experiment(table: Table, expids: list[str], limits: dict[str, date | None], limit: str) -> None:
    """
    Checks the Date of each row of a table against a date of its experiment: limits holds the StartDate of each
    experiment if limit is "StartDate", which no Date can be before, or the EndDate if it is "EndDate".
    """
    days = table.dates("Date")
    bounds = [limits[expid] for expid in expids]

    def fits(day: date, bound: date) -> bool:
        return day >= bound if limit == "StartDate" else day <= bound

    wrong = pd.Series([day is not None and bound is not None and not fits(day, bound)
                       for day, bound in zip(days, bounds)], dtype=bool)
    fail_first(table, "Date", wrong, lambda i: (
        f"'{days[i]}' is {'before' if limit == 'StartDate' else 'after'} {bounds[i]}, the {limit} of experiment "
        f"'{expids[i]}'." + swap_hint(days[i], lambda day: fits(day, bounds[i]))))


def dates_within_experiment(tables: dict[str, Table]) -> None:
    """
    Phenotype dates are not before their experiment's StartDate, and Management and Climate dates not after its
    EndDate. The other way round is normal: seed is scored after harvest, and fertiliser given before sowing.
    """
    experiments = tables["Experiment"]
    starts = dict(zip(experiments.data["EXPID"], experiments.dates("StartDate")))
    ends = dict(zip(experiments.data["EXPID"], experiments.dates("EndDate")))
    units = tables["ExperimentalUnit"]
    unit_experiment = dict(zip(units.data["UNID"], units.data["EXPID"]))
    phenotype = tables["Phenotype"]
    date_within_experiment(phenotype, [unit_experiment[unid] for unid in phenotype.data["UNID"]], starts,
                           "StartDate")
    for name in ("Management", "Climate"):
        table = tables.get(name)
        if table is not None:
            date_within_experiment(table, list(table.data["EXPID"]), ends, "EndDate")


def seed_lot_harvested_before_scoring(tables: dict[str, Table]) -> None:
    """A unit is scored after the seed lot it was sown from was harvested."""
    seed_lots = tables["SeedLot"]
    harvested = dict(zip(seed_lots.data["SLID"], seed_lots.dates("HarvestDate")))
    units = tables["ExperimentalUnit"]
    unit_seed_lot = dict(zip(units.data["UNID"], units.data["SLID"]))
    phenotype = tables["Phenotype"]
    unids = list(phenotype.data["UNID"])
    slids = [unit_seed_lot[unid] for unid in unids]
    days = phenotype.dates("Date")
    wrong = pd.Series([day is not None and harvested[slid] is not None and day <= harvested[slid]
                       for day, slid in zip(days, slids)], dtype=bool)
    fail_first(phenotype, "Date", wrong, lambda i: (
        f"'{days[i]}' is not after {harvested[slids[i]]}, the HarvestDate of seed lot '{slids[i]}', which unit "
        f"'{unids[i]}' was sown from." + swap_hint(days[i], lambda day: day > harvested[slids[i]])))


RULES: dict[str, Callable[[dict[str, Table]], None]] = {
    "climate_unit_in_experiment": climate_unit_in_experiment,
    "one_intervention_per_unit": one_intervention_per_unit,
    "every_experiment_has_a_unit": every_experiment_has_a_unit,
    "no_seed_lot_descends_from_itself": no_seed_lot_descends_from_itself,
    "parent_seed_lot_harvested_first": parent_seed_lot_harvested_first,
    "experiment_starts_before_it_ends": experiment_starts_before_it_ends,
    "dates_within_experiment": dates_within_experiment,
    "seed_lot_harvested_before_scoring": seed_lot_harvested_before_scoring,
}


# ---------------------------------------------------------------- Validation

def validate(raws: dict[str, RawTable], schema: Schema, notice: Callable[[str], None] = lambda message: None,
             warning: Callable[[str], None] = lambda message: None) -> dict[str, Table]:
    """
    Checks the tables against the schema, in this order, and stops at the first error:
    which tables there are, their columns and rows, the values in each column, the Features table,
    the links between tables, the Features cells, and the rules that span tables.
    notice(message) is called for what is good to know, and warning(message) for what may be a mistake.
    Returns:
        dict: table name -> checked table, in the schema's order.
    Raises:
        ValidationError: At the first error, saying where it is and how to fix it.
    """
    for definition in schema.tables.values():
        if definition.name not in raws:
            if definition.required:
                raise ValidationError(f"The {definition.name} table is missing. {definition.add_hint()}")
            notice(f"There is no {definition.name} table. It is optional, so it was skipped.")

    tables = {name: build_table(definition, raws[name], schema)
              for name, definition in schema.tables.items() if name in raws}
    for table in tables.values():
        if table.definition.required and table.data.empty:
            raise ValidationError(f"{table.source} has no rows, but {table.name} needs at least one.")

    for table in tables.values():
        check_cells(table, warning)
    features = check_feature_definitions(tables.get("Features"))
    for table in tables.values():
        for column in table.definition.columns.values():
            if column.links_to:
                check_links(table, column, tables, schema)
    for table in tables.values():
        for column in table.definition.columns.values():
            if column.type == "features":
                check_feature_cells(table, column.name, features, schema)
    for rule in schema.rules:
        RULES[rule](tables)
    return tables


def validate_files(files: list[tuple[str, Path]], schema: Schema | None = None,
                   notice: Callable[[str], None] = lambda message: None,
                   warning: Callable[[str], None] = lambda message: None) -> dict[str, Table]:
    """Reads the tables from CSV files or one collection sheet, and checks them. See `load_files` and `validate`."""
    schema = schema or load_schema()
    return validate(load_files(files, schema, notice, warning), schema, notice, warning)


def summary(tables: dict[str, Table], schema: Schema) -> list[str]:
    """Lines saying that the tables passed, and how many rows each has."""
    return ([f"Validation successful! All {len(tables)} tables pass schema version {schema.version}."]
            + [f"{table.name}: {len(table.data)} row{'' if len(table.data) == 1 else 's'}, from {table.source}"
               for table in tables.values()])


def write_tables(tables: dict[str, Table], folder: Path) -> None:
    """Writes each table to <folder>/<table>.csv, in UTF-8."""
    folder.mkdir(parents=True, exist_ok=True)
    for name, table in tables.items():
        (folder / f"{name}.csv").write_text(table.to_csv(), encoding="utf-8", newline="")


def tested_versions(path: Path = REQUIREMENTS_FILE) -> dict[str, str]:
    """Package -> version, from requirements.txt: the versions that the tests and the browser app run with."""
    versions = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip() and not line.startswith("#"):
            package, _, version = line.strip().partition("==")
            versions[package] = version
    return versions


def version_notices(path: Path = REQUIREMENTS_FILE) -> list[str]:
    """
    One line for each package installed in another version than the one in requirements.txt, as results may then
    differ from the tests'. Packages that are not installed are not used, so they are left out.
    """
    notices = []
    for package, tested in tested_versions(path).items():
        try:
            installed = metadata.version(package)
        except metadata.PackageNotFoundError:
            continue
        if installed != tested:
            notices.append(f"This validator is tested with {package} {tested}, but {package} {installed} "
                           "is installed, so results may differ.")
    return notices


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Check phenotype database tables against the schema. "
                                                 "Stops at the first error.")
    parser.add_argument("paths", nargs="+", type=Path,
                        help="CSV files named after their tables, a folder of them, or one collection sheet (.xlsx)")
    parser.add_argument("--out", type=Path,
                        help="if the tables pass, write them to this folder as CSV files "
                             "(UTF-8, comma-separated, dates YYYY-MM-DD)")
    args = parser.parse_args(argv)

    files = []
    for path in args.paths:
        if path.is_dir():
            files += [(p.name, p) for p in sorted(path.iterdir())
                      if p.suffix.lower() in (".csv", ".xlsx") and not p.name.startswith("~$")]
        elif path.is_file():
            files.append((path.name, path))
        else:
            parser.error(f"{path} does not exist")
    if args.out and any(path.parent.resolve() == args.out.resolve() for _, path in files):
        parser.error("--out has to be a different folder than the one the tables are read from")

    schema = load_schema()
    print(f"Checking against schema version {schema.version}")
    for notice in version_notices():
        print(f"Notice: {notice}")
    try:
        tables = validate_files(files, schema, notice=lambda message: print(f"Notice: {message}"),
                                warning=lambda message: print(f"Warning: {message}"))
    except ValidationError as e:
        print(f"Validation failed: {e}")
        return 1
    print("\n".join(summary(tables, schema)))
    if args.out:
        write_tables(tables, args.out)
        print(f"Wrote the tables to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
